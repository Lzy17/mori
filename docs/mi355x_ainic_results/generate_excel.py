#!/usr/bin/env python3
"""Generate Excel workbook from MORI benchmark results."""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

wb = openpyxl.Workbook()
ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

# ---- Sheet 1: environment ----
ws = wb.active
ws.title = "environment"
headers = [
    "timestamp", "hostname", "docker_image_container", "git_sha",
    "rocm_version", "torch_version", "python_version", "gpu_model",
    "gpu_count", "nic_type", "rdma_devices", "tcp_bootstrap_interface", "notes"
]
ws.append(headers)
ws.append([
    ts,
    "smci355-ccs-aus-g12-06",
    "rocm/sgl-dev:v0.5.10rc0-rocm700-mi35x-20260331 / mori_bench",
    "13ec475c (installed) / 241461c0 (latest)",
    "7.0.0",
    "2.9.0a0+git7bcbafe (HIP 7.0.51831-a3e329ad8)",
    "3.10.12",
    "AMD Instinct MI355X",
    8,
    "AINIC (Pollara / ionic)",
    "ionic_0..ionic_8 (9 total, ionic_6=mgmt MTU1500, ionic_2/3 DOWN on node2)",
    "enp196s0 (GLOO_SOCKET_IFNAME)",
    "Node 0: 10.235.24.254, Node 1: 10.235.58.242"
])
for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# ---- Sheet 2: mori_ep_intranode ----
ws2 = wb.create_sheet("mori_ep_intranode")
headers2 = [
    "timestamp", "command", "status", "duration_sec", "num_gpus",
    "hidden_dim", "num_tokens", "topk", "dtype", "output_file",
    "raw_stdout_summary", "key_metrics", "notes"
]
ws2.append(headers2)

# Parse intra-node results
dispatch_bws = [357.0, 371.25, 368.5, 367.125, 365.875, 373.625, 372.5, 368.5, 371.875, 369.75]
combine_bws = [392.375, 395.375, 396.75, 395.375, 394.5, 396.75, 396.25, 396.75, 393.375, 394.0]
avg_dispatch = sum(dispatch_bws) / len(dispatch_bws)
avg_combine = sum(combine_bws) / len(combine_bws)

ws2.append([
    ts,
    "python3 tests/python/ops/bench_dispatch_combine.py",
    "SUCCESS",
    "~30",
    8,
    7168,
    4096,
    8,
    "bf16",
    "results/logs/ep_intranode.log",
    f"Dispatch avg={avg_dispatch:.1f} GB/s, Combine avg={avg_combine:.1f} GB/s",
    f"Dispatch BW: {avg_dispatch:.1f} GB/s, Combine BW: {avg_combine:.1f} GB/s, ~296 MB/round",
    "No env var changes needed for intra-node. Default config used."
])
for cell in ws2[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# ---- Sheet 3: mori_ep_internode ----
ws3 = wb.create_sheet("mori_ep_internode")
headers3 = [
    "timestamp", "command", "status", "nnodes", "node_rank",
    "nproc_per_node", "master_addr", "master_port", "gloo_socket_ifname",
    "mori_rdma_devices", "output_file", "total_tokens", "rdma_tokens",
    "total_bw", "derived_rdma_only_bw", "notes"
]
ws3.append(headers3)
ws3.append([
    ts,
    ("torchrun --nnodes=2 --node_rank=0 --nproc_per_node=1 "
     "--master_addr=10.235.24.254 --master_port=1241 "
     "examples/ops/dispatch_combine/test_dispatch_combine_internode.py --cmd bench --num-qp 1"),
    "SUCCESS (with --num-qp 1, single NIC)",
    2, 0, 1, "10.235.24.254", 1241, "enp196s0",
    "ionic_0",
    "results/logs/ep_internode_success.log",
    4096, 8192, "16 GB/s (total), 4 GB/s (RDMA)", "4 GB/s",
    ("Succeeded with --num-qp 1 and MORI_RDMA_DEVICES=ionic_0 (single NIC). "
     "MORI_DISABLE_TOPO=1 required. "
     "Dispatch: avg 27459 us, RDMA 4 GB/s, total 16 GB/s. "
     "Combine: avg 26711 us, RDMA 4 GB/s, total 17 GB/s. "
     "Hangs with 6 NICs or default num-qp (4). "
     "Root cause: ionic QP scaling issue with many concurrent QPs.")
])
for cell in ws3[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# ---- Sheet 4: mori_io ----
ws4 = wb.create_sheet("mori_io")
headers4 = [
    "timestamp", "command", "status", "nnodes", "node_rank",
    "nproc_per_node", "host", "buffer_size", "transfer_batch_size",
    "enable_batch_transfer", "enable_sess", "output_file",
    "max_bw_gbps", "avg_bw_gbps", "min_lat_us", "avg_lat_us", "notes"
]
ws4.append(headers4)

# IO benchmark data points from the actual output
io_data = [
    (8, 256, 0.07, 0.07, 28.37, 29.15),
    (16, 256, 0.14, 0.14, 28.61, 29.25),
    (32, 256, 0.28, 0.28, 28.85, 29.77),
    (64, 256, 0.56, 0.55, 29.09, 30.01),
    (128, 256, 1.12, 1.07, 29.33, 30.76),
    (256, 256, 2.15, 2.05, 30.52, 31.94),
    (512, 256, 4.01, 3.74, 32.66, 35.01),
    (1024, 256, 7.09, 6.73, 36.95, 38.94),
    (2048, 256, 11.45, 10.75, 45.78, 48.76),
    (4096, 256, 16.41, 15.63, 63.90, 67.08),
    (8192, 256, 21.20, 19.97, 98.94, 104.99),
    (16384, 256, 24.60, 23.62, 170.47, 177.56),
    (32768, 256, 26.61, 26.10, 315.19, 321.37),
    (65536, 256, 27.88, 27.43, 601.77, 611.58),
    (131072, 256, 28.44, 28.19, 1179.93, 1190.30),
    (262144, 256, 28.80, 28.61, 2330.54, 2345.73),
    (524288, 256, 28.95, 28.78, 4635.81, 4663.33),
    (1048576, 256, 28.93, 28.81, 9278.77, 9315.96),
]

cmd_io = ("torchrun --nnodes=2 --node_rank=0 --nproc_per_node=1 "
          "--master_addr=10.235.24.254 --master_port=1243 "
          "tests/python/io/benchmark.py --host=<node_own_ip> "
          "--all --enable-sess --enable-batch-transfer")

for row in io_data:
    msg_size, batch, max_bw, avg_bw, min_lat, avg_lat = row
    ws4.append([
        ts, cmd_io, "SUCCESS", 2, 0, 1,
        "10.235.24.254 (node0) / 10.235.58.242 (node1)",
        msg_size, batch, True, True,
        "results/logs/io_node0.log",
        max_bw, avg_bw, min_lat, avg_lat,
        f"RDMA read, 1 initiator + 1 target, ionic_0 (auto-selected)"
    ])

for cell in ws4[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# ---- Sheet 4b: mori_io_6nic ----
ws4b = wb.create_sheet("mori_io_6nic")
headers4b = [
    "timestamp", "command", "status", "nnodes",
    "num_initiator_dev", "num_target_dev", "num_qp_per_transfer",
    "transfer_batch_size", "enable_batch_transfer", "enable_sess",
    "initiator_rank", "nic_device",
    "msg_size_bytes", "max_bw_gbps", "avg_bw_gbps", "min_lat_us", "avg_lat_us",
    "notes"
]
ws4b.append(headers4b)

# 6-NIC IO data: per-rank results at key message sizes
io_6nic_data = {
    # rank: [(msg_size, max_bw, avg_bw, min_lat, avg_lat), ...]
    0: [
        (8, 0.07, 0.07, 29.33, 30.22), (64, 0.55, 0.53, 30.04, 31.15),
        (1024, 6.87, 6.53, 38.15, 40.14), (8192, 20.22, 19.69, 103.71, 106.51),
        (65536, 27.30, 26.84, 614.64, 625.04), (262144, 28.24, 28.03, 2376.32, 2394.21),
        (524288, 28.32, 28.22, 4739.52, 4756.38), (1048576, 28.30, 28.24, 9484.77, 9506.41),
    ],
    1: [
        (8, 0.07, 0.07, 28.61, 29.68), (64, 0.56, 0.53, 29.33, 31.04),
        (1024, 7.05, 6.59, 37.19, 39.76), (8192, 21.09, 20.06, 99.42, 104.55),
        (65536, 27.46, 27.06, 611.07, 620.04), (262144, 28.32, 28.17, 2369.40, 2382.50),
        (524288, 28.50, 28.40, 4708.77, 4726.00), (1048576, 28.65, 28.58, 9369.37, 9391.78),
    ],
    2: [
        (8, 0.07, 0.07, 29.80, 30.83), (64, 0.54, 0.52, 30.52, 31.77),
        (1024, 6.07, 5.83, 43.15, 44.99), (8192, 14.28, 13.95, 146.87, 150.32),
        (65536, 17.18, 17.10, 976.32, 981.19), (262144, 17.59, 17.53, 3814.46, 3827.18),
        (524288, 17.66, 17.61, 7599.83, 7621.64), (1048576, 17.74, 17.48, 15134.10, 15354.15),
    ],
    3: [
        (8, 0.07, 0.06, 30.28, 31.54), (64, 0.54, 0.51, 30.52, 32.01),
        (1024, 6.04, 5.74, 43.39, 45.69), (8192, 14.30, 13.97, 146.63, 150.16),
        (65536, 17.22, 17.10, 974.18, 981.05), (262144, 17.62, 17.54, 3808.98, 3825.42),
        (524288, 17.69, 17.62, 7586.96, 7617.64), (1048576, 17.74, 17.49, 15132.67, 15347.02),
    ],
    4: [
        (8, 0.07, 0.07, 30.04, 31.19), (64, 0.54, 0.51, 30.52, 31.97),
        (1024, 6.18, 5.87, 42.44, 44.65), (8192, 14.93, 14.57, 140.43, 143.96),
        (65536, 18.20, 18.06, 921.73, 929.12), (262144, 18.60, 18.55, 3608.23, 3618.52),
        (524288, 18.68, 18.64, 7185.22, 7201.06), (1048576, 18.65, 18.60, 14392.38, 14428.44),
    ],
    5: [
        (8, 0.07, 0.07, 28.85, 29.60), (64, 0.55, 0.53, 29.56, 30.76),
        (1024, 6.32, 6.06, 41.48, 43.29), (8192, 15.38, 14.74, 136.38, 142.29),
        (65536, 18.30, 18.15, 916.96, 924.29), (262144, 18.72, 18.69, 3584.86, 3591.54),
        (524288, 18.79, 18.77, 7144.45, 7151.08), (1048576, 18.89, 18.87, 14210.46, 14225.72),
    ],
}

nic_names = {0: "ionic_0", 1: "ionic_1", 2: "ionic_4", 3: "ionic_5", 4: "ionic_7", 5: "ionic_8"}

cmd_io_6nic = ("torchrun --nnodes=2 --node_rank=0 --nproc_per_node=1 "
               "--master_addr=10.235.24.254 --master_port=1245 "
               "tests/python/io/benchmark.py --host=<node_own_ip> "
               "--all --enable-sess --enable-batch-transfer "
               "--num-initiator-dev 6 --num-target-dev 6")

for rank, data_points in io_6nic_data.items():
    for msg_size, max_bw, avg_bw, min_lat, avg_lat in data_points:
        ws4b.append([
            ts, cmd_io_6nic, "SUCCESS", 2,
            6, 6, 1, 256, True, True,
            rank, nic_names[rank],
            msg_size, max_bw, avg_bw, min_lat, avg_lat,
            "RDMA read, 6 initiator + 6 target, session+batch"
        ])

for cell in ws4b[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# ---- Sheet 5: failures_and_fixes ----
ws5 = wb.create_sheet("failures_and_fixes")
headers5 = [
    "timestamp", "component", "attempted_command", "failure_symptom",
    "likely_root_cause", "fix_applied", "final_status", "notes"
]
ws5.append(headers5)

failures = [
    [
        ts, "MORI-EP Inter-node",
        "torchrun --nnodes=2 --nproc_per_node=8 ... --bench",
        "unrecognized arguments: --bench",
        "Docs say --bench but script expects --cmd bench",
        "Changed to --cmd bench",
        "FIXED",
        "Doc/script mismatch"
    ],
    [
        ts, "MORI-EP Inter-node",
        "torchrun --nnodes=2 --nproc_per_node=8 ... --cmd bench",
        "Assertion rankInNode < 8 failed (128 total ranks created)",
        "Script internally spawns gpu_per_node(8) processes; torchrun nproc_per_node=8 creates 8 outer; 8*8=64 per node",
        "Changed to --nproc_per_node=1 (torchrun launches 1 per node, script spawns 8 internally)",
        "FIXED",
        "WORLD_SIZE from torchrun = num_nodes; script handles GPU parallelism"
    ],
    [
        ts, "MORI-EP Inter-node",
        "torchrun --nnodes=2 --nproc_per_node=1 ... (default topo)",
        "Assertion 'no rdma device found' on node 1",
        "Topology detection (MatchGpuAndNic) fails for ionic devices",
        "Set MORI_DISABLE_TOPO=1",
        "FIXED (topo issue), but led to next hang",
        "ionic not in topo detection code path"
    ],
    [
        ts, "MORI-EP Inter-node",
        "... MORI_DISABLE_TOPO=1 MORI_RDMA_DEVICES=^ionic_6",
        "Hang during RDMA endpoint AllToAll exchange (context.cpp:276) with 6 NICs",
        "ionic QP scaling issue: too many concurrent QPs (6 NICs x 16 ranks x default 4 QPs)",
        "Use --num-qp 1 and MORI_RDMA_DEVICES=ionic_0 (single NIC). Hangs resolved.",
        "FIXED (with single NIC + --num-qp 1)",
        "EP internode works with 1 NIC + 1 QP. Hangs with 6 NICs or default QP count. "
        "RDMA BW 4 GB/s, Total BW 16 GB/s dispatch / 17 GB/s combine."
    ],
    [
        ts, "MORI-IO",
        "torchrun ... benchmark.py ...",
        "ModuleNotFoundError: No module named 'prettytable'",
        "Missing Python dependency in Docker image",
        "pip install prettytable",
        "FIXED",
        "Should be in Docker image requirements"
    ],
    [
        ts, "MORI-IO",
        "torchrun ... benchmark.py --host=10.235.24.254",
        "syscall failed with Cannot assign requested address (on non-host node)",
        "--host IP used by ALL ranks for TCP bind; non-host node can't bind to remote IP",
        "Each node must use its OWN IP for --host (node0: 10.235.24.254, node1: 10.235.58.242)",
        "FIXED",
        "Doc is ambiguous about --host; should clarify it's per-node"
    ],
]

for f in failures:
    ws5.append(f)

for cell in ws5[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.font = Font(bold=True, color="FFFFFF")

# Auto-size columns
for ws in [wb.active, ws2, ws3, ws4, ws4b, ws5]:
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, min(len(str(cell.value)), 60))
        ws.column_dimensions[col_letter].width = max_length + 2

wb.save("/data/sglang_bench/mori_bench/results/mori_benchmark_results.xlsx")
print("Excel saved to /data/sglang_bench/mori_bench/results/mori_benchmark_results.xlsx")
